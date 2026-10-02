#!/usr/bin/env python3
"""
Generate the editable design & documentation file (.docx) from the same
design tokens that theme the Tableau workbook.

    python3 theme/make_doc.py

Output: Air_Quality_Dashboard_Documentation.docx (repo root)
Re-run after changing theme/palette.json to keep the document in sync.
"""
from __future__ import annotations

import json
from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parent.parent
T = json.loads((ROOT / "theme" / "palette.json").read_text(encoding="utf-8"))
C = T["color"]
TY = T["type"]
VERSION = T.get("version", "1.0.0")

def hx(token: str) -> str:
    return token.lstrip("#").lower()

def rgb(token: str) -> RGBColor:
    h = hx(token)
    return RGBColor(int(h[0:2], 16), int(h[2:4], 16), int(h[4:6], 16))

INK = C["ink"]["heading"]
BODY = C["ink"]["body"]
MUTED = C["ink"]["muted"]
BRAND = C["brand"]["primary"]

# ---------------------------------------------------------------- helpers --
def shade(cell, hexcolor: str):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:fill"), hx(hexcolor))
    tcPr.append(shd)

def set_cell(cell, text, *, bold=False, color=BODY, size=9.5, align=None):
    cell.text = ""
    p = cell.paragraphs[0]
    run = p.add_run(text)
    run.bold = bold
    run.font.size = Pt(size)
    run.font.color.rgb = rgb(color)
    run.font.name = "Calibri"
    if align:
        p.alignment = align
    return cell

def add_table(doc: Document, headers, rows, widths=None, swatch_col=None,
              header_fill=None):
    tbl = doc.add_table(rows=1, cols=len(headers))
    tbl.style = "Table Grid"
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, h in enumerate(headers):
        cell = tbl.rows[0].cells[i]
        set_cell(cell, h, bold=True, color="#FFFFFF", size=9.5)
        shade(cell, header_fill or INK)
    for r_i, row in enumerate(rows):
        cells = tbl.add_row().cells
        fill = "#FFFFFF" if r_i % 2 == 0 else hx(C["surface"]["canvas"])
        for c_i, value in enumerate(row):
            if swatch_col is not None and c_i == swatch_col:
                set_cell(cells[c_i], "", size=2)
                shade(cells[c_i], value if value.startswith("#") else "#FFFFFF")
            else:
                set_cell(cells[c_i], str(value))
            if c_i != swatch_col:
                shade(cells[c_i], fill)
    if widths:
        for row in tbl.rows:
            for i, w in enumerate(widths):
                row.cells[i].width = Inches(w)
    return tbl

def h(doc, text, level, color=INK):
    heading = doc.add_heading(text, level=level)
    for run in heading.runs:
        run.font.color.rgb = rgb(color)
        run.font.name = "Calibri"
    return heading

def p(doc, text, *, size=10.5, color=BODY, bold=False, italic=False,
      space_after=6, align=None):
    para = doc.add_paragraph()
    run = para.add_run(text)
    run.font.size = Pt(size)
    run.font.color.rgb = rgb(color)
    run.bold = bold
    run.italic = italic
    run.font.name = "Calibri"
    para.paragraph_format.space_after = Pt(space_after)
    if align:
        para.alignment = align
    return para

def bullet(doc, text, *, bold_prefix=None, size=10.5):
    para = doc.add_paragraph(style="List Bullet")
    if bold_prefix:
        r = para.add_run(bold_prefix)
        r.bold = True
        r.font.size = Pt(size)
        r.font.color.rgb = rgb(INK)
        r.font.name = "Calibri"
    r = para.add_run(text)
    r.font.size = Pt(size)
    r.font.color.rgb = rgb(BODY)
    r.font.name = "Calibri"
    return para

def swatch_row(group, token, hexval, applied):
    return [group, token, hexval.upper(), hexval, applied]

# ----------------------------------------------------------------- document --
doc = Document()

# page + base style
for section in doc.sections:
    section.left_margin = Inches(0.8)
    section.right_margin = Inches(0.8)
    section.top_margin = Inches(0.7)
    section.bottom_margin = Inches(0.7)
normal = doc.styles["Normal"]
normal.font.name = "Calibri"
normal.font.size = Pt(10.5)
normal.font.color.rgb = rgb(BODY)

# ---- cover block ------------------------------------------------------------
title = doc.add_paragraph()
title.alignment = WD_ALIGN_PARAGRAPH.LEFT
run = title.add_run("INDIA AIR QUALITY INTELLIGENCE")
run.font.size = Pt(26)
run.bold = True
run.font.color.rgb = rgb(INK)
run.font.name = "Calibri"

sub = doc.add_paragraph()
run = sub.add_run("Tableau Dashboard Suite  \u00b7  Design & Documentation")
run.font.size = Pt(14)
run.font.color.rgb = rgb(BRAND)
run.font.name = "Calibri"

meta = doc.add_paragraph()
run = meta.add_run(f"Atmos Design System v{VERSION}  \u00b7  {date.today():%d %B %Y}  \u00b7  "
                   f"Workbook: Air_Pollution_Dashboard.twbx")
run.font.size = Pt(9.5)
run.font.color.rgb = rgb(MUTED)
run.font.name = "Calibri"

# accent rule (single-row table, shaded)
rule = doc.add_table(rows=1, cols=1)
rule.rows[0].cells[0].text = ""
shade(rule.rows[0].cells[0], BRAND)
rule.rows[0].height = Pt(4)
doc.add_paragraph()

# ---- 1. overview ------------------------------------------------------------
h(doc, "1.  What was delivered", 1)
p(doc, "The original workbook contained 16 worksheets and two dashboards styled with "
       "ad-hoc colors (yellow panels, black title bars, a dark basemap, mixed fonts). "
       "This delivery rebuilds both dashboards on a professional card-based layout, "
       "introduces the Atmos modular color system, adds a city locator map, and ships a "
       "theme builder that can re-skin the entire workbook from a single token file.")
add_table(doc,
    ["Item", "Description"],
    [
        ["Air_Pollution_Dashboard.twbx", "The themed Tableau workbook \u2014 fully editable in Tableau Desktop. Two dashboards, 17 worksheets, all filters and parameters preserved."],
        ["theme/palette.json", "Design tokens \u2014 the single source of truth for every color and font size in the workbook and this document."],
        ["theme/apply_theme.py", "Theme builder: restyles worksheets, pins mark colors, rebuilds dashboards, generates button assets and design previews. Idempotent."],
        ["preview/*.png", "Design previews of both dashboards, rendered from the same layout spec and tokens."],
        ["DESIGN_SYSTEM.md", "Developer-facing reference version of this document."],
        ["This document", "Editable stakeholder documentation \u2014 regenerate any time with theme/make_doc.py."],
    ],
    widths=[2.1, 4.6])

h(doc, "2.  The data", 1)
p(doc, "Daily ambient air-quality readings published by the CPCB continuous monitoring "
       "network: 26 Indian cities, January 2021 \u2013 December 2022 (18,928 daily records). "
       "Each record carries an AQI index (1\u20135 scale; unhealthy \u2265 4) plus concentrations "
       "of eight pollutants \u2014 CO, NO, NO\u2082, O\u2083, SO\u2082, PM2.5, PM10 and NH\u2083 \u2014 benchmarked "
       "against Indian NAAQS annual and 24-hour limits.")

# ---- 3. dashboards ----------------------------------------------------------
h(doc, "3.  Dashboard guide", 1)

h(doc, "3.1  Nationwide Overview", 2)
p(doc, "The executive landing view: national KPIs, a city-level AQI map, and the two "
       "signature analytical charts. A Year filter in the header compares 2021 vs 2022 "
       "across every chart; the header button drills into city-level detail.")
doc.add_picture(str(ROOT / "preview" / "Nationwide_Overview.png"), width=Inches(6.9))
doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
add_table(doc,
    ["Band", "Contents"],
    [
        ["Header", "Title + subtitle, Year filter (dropdown), City Drilldown navigation button, brand accent rule"],
        ["KPI band", "National Average AQI \u00b7 YoY Change in AQI \u00b7 Most Polluted City \u00b7 Cleanest City \u00b7 Peak Pollution Month"],
        ["Chart grid", "Average AQI by City (bubble map, color = category, size = avg AQI) \u00b7 Monthly AQI Trend by Year (matrix) \u00b7 Top 5 Air Pollutants by Avg Concentration (bars)"],
        ["Footnote band", "Records Monitored stat card + \u201cHow to read\u201d guidance"],
    ],
    widths=[1.3, 5.4])

h(doc, "3.2  City Drilldown", 2)
p(doc, "The analytical deep-dive. A control strip drives every chart: pick a city and a "
       "pollutant and the whole page updates \u2014 including a locator map that zooms to the "
       "selected city. A header button returns to the national overview.")
doc.add_picture(str(ROOT / "preview" / "City_Drilldown.png"), width=Inches(6.9))
doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
add_table(doc,
    ["Band", "Contents"],
    [
        ["Header", "Title + subtitle, Nationwide View navigation button, brand accent rule"],
        ["Control strip", "City filter (dropdown) \u00b7 Pollutant parameter \u00b7 Year color legend"],
        ["KPI band", "Selected City \u00b7 Avg AQI \u00b7 Peak Pollution Month \u00b7 % Unhealthy AQI Days"],
        ["Chart grid", "Average concentration of the selected pollutant, 2021 vs 2022, against the NAAQS annual limit \u00b7 Selected City Location (locator map) \u00b7 Monthly concentration trend by year \u00b7 % of days above the NAAQS 24-hour limit"],
        ["Footnote band", "\u201cHow to read\u201d guidance + source note"],
    ],
    widths=[1.3, 5.4])

# ---- 4. design system -------------------------------------------------------
h(doc, "4.  The Atmos design system", 1)
p(doc, "Every color used anywhere \u2014 dashboard chrome, worksheet surfaces, mark colors, "
       "fonts, buttons, and this document \u2014 resolves to a token in theme/palette.json. "
       "Discrete mark colors are additionally pinned in the workbook XML as value \u2192 color "
       "maps, so the same value always renders the same color on every sheet.")

h(doc, "4.1  Surfaces, ink and lines", 2)
rows = [
    swatch_row("Surface", "canvas", C["surface"]["canvas"], "Dashboard background"),
    swatch_row("Surface", "card", C["surface"]["card"], "All cards; worksheet table and title backgrounds"),
    swatch_row("Surface", "header", C["surface"]["header"], "Dashboard header band"),
    swatch_row("Ink", "heading", C["ink"]["heading"], "Chart titles, KPI values, emphasis text"),
    swatch_row("Ink", "body", C["ink"]["body"], "Worksheet text, footnotes"),
    swatch_row("Ink", "muted", C["ink"]["muted"], "KPI labels, captions, reference lines, tooltips"),
    swatch_row("Ink", "inverse", C["ink"]["inverse"], "Title text on the header"),
    swatch_row("Ink", "inverseMuted", C["ink"]["inverseMuted"], "Subtitle text on the header"),
    swatch_row("Line", "hairline", C["line"]["hairline"], "Card borders"),
    swatch_row("Line", "headerRule", C["line"]["headerRule"], "Navigation button border"),
]
add_table(doc, ["Group", "Token", "Hex", "", "Applied to"], rows,
          widths=[0.9, 1.25, 0.95, 0.55, 3.05], swatch_col=3)

h(doc, "4.2  Brand, series and status", 2)
rows = [
    swatch_row("Brand", "primary", C["brand"]["primary"], "Accent rule, KPI card accent strips, Avg-AQI marks, 2022 series, parameter highlight"),
    swatch_row("Brand", "primaryDark", C["brand"]["primaryDark"], "Reserved for emphasis / hover"),
    swatch_row("Series", "prior (2021)", C["series"]["prior"], "Prior-year marks and bars"),
    swatch_row("Series", "current (2022)", C["series"]["current"], "Current-year marks and bars"),
    swatch_row("Status", "good", C["status"]["good"], "Low AQI category; within NAAQS limits"),
    swatch_row("Status", "moderate", C["status"]["moderate"], "Moderate AQI category"),
    swatch_row("Status", "bad", C["status"]["bad"], "High AQI category; exceedance days"),
]
add_table(doc, ["Group", "Token", "Hex", "", "Applied to"], rows,
          widths=[0.9, 1.25, 0.95, 0.55, 3.05], swatch_col=3)

h(doc, "4.3  Categorical and pollutant palette", 2)
pollutants = [
    ("PM2.5", "pm2_5"), ("PM10", "pm10"), ("CO", "co"), ("NO2", "no2"),
    ("O3", "o3"), ("SO2", "so2"), ("NH3", "nh3"), ("NO", "no"),
]
rows = [swatch_row("Pollutant", name, C["pollutant"][key],
                   f"Fixed color for {name} \u2014 bars, legends, encodings")
        for name, key in pollutants]
rows += [swatch_row("Categorical", key, val, "General-purpose distinct series color")
         for key, val in C["categorical"].items()
         if val.lower() not in {c.lower() for c in C["pollutant"].values()}]
add_table(doc, ["Group", "Series", "Hex", "", "Applied to"], rows,
          widths=[0.9, 1.25, 0.95, 0.55, 3.05], swatch_col=3)

h(doc, "4.4  Semantic color locks", 2)
bullet(doc, "2021 \u2192 series.prior, 2022 \u2192 series.current \u2014 every year-colored mark in the workbook.", bold_prefix="Years:  ")
bullet(doc, "Low \u2192 good, Moderate \u2192 moderate, High \u2192 bad (map bubbles and legends).", bold_prefix="AQI category:  ")
bullet(doc, "Normal \u2192 good, High \u2192 bad (days above the NAAQS 24-hour limit).", bold_prefix="Exceedance flag:  ")
bullet(doc, "each of the eight pollutants keeps one fixed color across all sheets.", bold_prefix="Pollutants:  ")

h(doc, "4.5  Typography", 2)
add_table(doc,
    ["Role", "Font", "Size", "Weight"],
    [
        ["Dashboard title", "Tableau Bold", str(TY["dashboardTitle"]), "Bold"],
        ["Dashboard subtitle", "Tableau", str(TY["dashboardSubtitle"]), "Regular"],
        ["Chart titles", "Tableau Bold", str(TY["chartTitle"]), "Bold"],
        ["KPI labels", "Tableau", str(TY["kpiLabel"]), "Regular (uppercase)"],
        ["KPI values", "Tableau Bold", f"{TY['kpiValueText']}\u2013{TY['kpiValueNumeric']}", "Bold"],
        ["Footnotes", "Tableau", str(TY["footnote"]), "Regular"],
    ],
    widths=[1.7, 1.5, 1.2, 1.6])
p(doc, "Tableau\u2019s bundled font family is used throughout so the workbook renders "
       "identically on every machine \u2014 no font installation required.", italic=True, size=9.5)

h(doc, "4.6  Layout grid", 2)
bullet(doc, "1% page margins and 1% gutters between cards on both dashboards.")
bullet(doc, "Full-bleed navy header band (9.5% height) followed by a brand accent rule.")
bullet(doc, "Bands stack vertically: header \u2192 controls (drilldown) \u2192 KPI band \u2192 chart grid \u2192 footnote.")
bullet(doc, "Sizing mode is Automatic, so the percent-based grid adapts to any screen.")

# ---- 5. worksheets ----------------------------------------------------------
h(doc, "5.  Worksheet reference", 1)
ws_rows = [
    ["National AQI ", "KPI", "National average AQI (Year filter aware)", "Nationwide \u00b7 KPI band"],
    ["% Change", "KPI", "Year-over-year change in average AQI", "Nationwide \u00b7 KPI band"],
    ["Most Polluted City", "KPI", "City with the highest average AQI", "Nationwide \u00b7 KPI band"],
    ["Cleanest City", "KPI", "City with the lowest average AQI", "Nationwide \u00b7 KPI band"],
    ["Peak Pollution Month", "KPI", "Month with the highest average AQI", "Nationwide \u00b7 KPI band"],
    ["Total Records Monitored", "KPI", "Count of daily readings in scope", "Nationwide \u00b7 footnote"],
    ["Map", "Map", "Bubble map of average AQI for all 26 cities", "Nationwide \u00b7 chart grid"],
    ["Monthly Y Variation", "Matrix", "Monthly average AQI for 2021, 2022 and the YoY delta", "Nationwide \u00b7 chart grid"],
    ["Top 5 Pollutants", "Bar chart", "Top five pollutants by average concentration", "Nationwide \u00b7 chart grid"],
    ["City Selector", "KPI", "Name of the city currently selected by the filter", "Drilldown \u00b7 KPI band"],
    ["Avg AQI ", "KPI", "Average AQI for the selected city", "Drilldown \u00b7 KPI band"],
    ["C Peak Pollution Month", "KPI", "Peak pollution month for the selected city", "Drilldown \u00b7 KPI band"],
    ["Unhealthy AQI Days", "KPI", "Share of days with AQI \u2265 4 for the selected city", "Drilldown \u00b7 KPI band"],
    ["Avg Conc Year", "Bar chart", "Selected pollutant, 2021 vs 2022, vs NAAQS annual limit", "Drilldown \u00b7 chart grid"],
    ["City Map", "Map", "Locator map that zooms to the selected city", "Drilldown \u00b7 chart grid"],
    ["Monthly Avg Concentration", "Bar chart", "Monthly concentration of the selected pollutant by year", "Drilldown \u00b7 chart grid"],
    ["Percentage of Days", "Bar chart", "Share of days above the NAAQS 24-hour limit, with threshold reference", "Drilldown \u00b7 chart grid"],
]
add_table(doc, ["Worksheet", "Type", "Shows", "Used on"], ws_rows,
          widths=[1.55, 0.85, 2.9, 1.4])

# ---- 6. theming -------------------------------------------------------------
h(doc, "6.  Re-theming the workbook", 1)
p(doc, "The color code is modular: change tokens in theme/palette.json, re-run one "
       "command, and the entire workbook \u2014 dashboards, worksheets, mark colors, map, "
       "buttons, previews and this document \u2014 re-skins together. The build is idempotent, "
       "so it can be run on an already-themed workbook at any time.")
for step in [
    "Edit theme/palette.json (colors and/or font sizes).",
    "Run: python3 theme/apply_theme.py",
    "Optional: python3 theme/make_doc.py to regenerate this document.",
    "Open Air_Pollution_Dashboard.twbx in Tableau \u2014 the new theme is applied.",
]:
    bullet(doc, step, bold_prefix=f"Step {['Edit', 'Run', 'Optional', 'Open'][ ['Edit','Run','Optional','Open'].index(step.split()[0]) ]}:  ") if False else bullet(doc, step)
p(doc, "Example \u2014 a warmer brand in two tokens:", bold=True, size=10)
warm = doc.add_paragraph()
run = warm.add_run('"brand": { "primary": "#B3541E", "primaryDark": "#8C3F16" },\n'
                   '"surface": { "canvas": "#F7F2EC", "card": "#FFFFFF", "header": "#3B2A1D" }')
run.font.name = "Consolas"
run.font.size = Pt(9.5)
run.font.color.rgb = rgb(INK)

# ---- 7. QA ------------------------------------------------------------------
h(doc, "7.  Quality assurance", 1)
for item in [
    "Generated workbook XML is well-formed and opens cleanly in Tableau.",
    "Zone geometry audited programmatically \u2014 no overlaps, nothing out of bounds on either dashboard.",
    "Every color literal in the workbook resolves to a palette token (zero stray colors).",
    "Data extracts (.hyper) byte-identical to the original upload \u2014 no data was touched.",
    "All navigation buttons resolve to their target windows; all worksheet references valid.",
    "Theme builder verified idempotent \u2014 consecutive runs produce identical workbook content.",
]:
    bullet(doc, item)

p(doc, "")
note = doc.add_paragraph()
run = note.add_run("Generated from theme/palette.json \u00b7 Atmos Design System "
                   f"v{VERSION} \u00b7 {date.today():%d %b %Y}")
run.font.size = Pt(8.5)
run.font.color.rgb = rgb(MUTED)
run.italic = True
note.alignment = WD_ALIGN_PARAGRAPH.CENTER

out = ROOT / "Air_Quality_Dashboard_Documentation.docx"
doc.save(out)
print(f"  \u2713 wrote {out.relative_to(ROOT)} ({out.stat().st_size/1024:.0f} KB)")
